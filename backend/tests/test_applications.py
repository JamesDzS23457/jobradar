from datetime import UTC, datetime, timedelta
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.core.db import SessionLocal
from app.main import app
from app.models import Job, Match


def _prepare_greenhouse_job(client: TestClient) -> int:
    client.post("/api/demo/reset")
    jobs = client.get("/api/jobs").json()["items"]
    applications = client.get("/api/applications").json()["items"]
    applied_ids = {item["job_id"] for item in applications}
    job = next(item for item in jobs if item["id"] not in applied_ids)
    with SessionLocal() as db:
        record = db.get(Job, job["id"])
        assert record is not None
        record.url = "https://boards.greenhouse.io/acme/jobs/123"
        record.posted_at = datetime.now(UTC)
        db.commit()
    return job["id"]


def test_application_preview_session_cv_and_submit():
    with TestClient(app) as client:
        job_id = _prepare_greenhouse_job(client)

        preview = client.get(f"/api/jobs/{job_id}/application?cv_source=tailored")
        assert preview.status_code == 200
        payload = preview.json()["preview"]
        assert payload["portal"]["mode"] == "auto_submit"
        assert payload["fields"]["email"]
        assert payload["can_auto_submit"] is True

        prepared = client.post(f"/api/jobs/{job_id}/application-session", json={"cv_source": "tailored"})
        assert prepared.status_code == 200
        session = prepared.json()["session"]
        assert session["status"] == "awaiting_review"
        assert session["cv_source"] == "tailored"

        session_detail = client.get(f"/api/application-sessions/{session['id']}")
        assert session_detail.status_code == 200
        assert session_detail.json()["session"]["job"]["id"] == job_id

        cv = client.get(session["cv_url"])
        assert cv.status_code == 200
        assert cv.content[:4] == b"%PDF"

        submitted = client.post(
            f"/api/extension/sessions/{session['id']}/status",
            json={"status": "submitted", "message": "Test submission"},
        )
        assert submitted.status_code == 200
        applications = client.get("/api/applications").json()["items"]
        assert any(item["job_id"] == job_id for item in applications)
        assert any(item["job_id"] == job_id and item["cv_source"] == "tailored" for item in applications)

        client.post("/api/demo/reset")


def test_auto_run_queues_allowlisted_job_and_reports_submission():
    with TestClient(app) as client:
        job_id = _prepare_greenhouse_job(client)
        prefs = client.put(
            "/api/preferences",
            json={"auto_submit_enabled": True, "auto_submit_max_per_run": 1, "freshness_days": 2},
        )
        assert prefs.status_code == 200

        run_response = client.post("/api/applications/auto-run", json={"max_items": 1})
        assert run_response.status_code == 200
        run = run_response.json()["run"]
        assert run["total"] == 1
        assert run["items"][0]["job_id"] == job_id
        assert run["items"][0]["status"] == "queued"
        session_id = run["items"][0]["session_id"]
        assert session_id

        claimed = client.get("/api/extension/next")
        assert claimed.status_code == 200
        claimed_session = claimed.json()["session"]
        assert claimed_session["id"] == session_id
        assert claimed_session["job"]["url"].startswith("https://boards.greenhouse.io/")

        finished = client.post(
            f"/api/extension/sessions/{session_id}/status",
            json={"status": "submitted", "message": "Auto-run test"},
        )
        assert finished.status_code == 200

        run_detail = client.get(f"/api/application-runs/{run['id']}")
        assert run_detail.status_code == 200
        result = run_detail.json()["run"]
        assert result["submitted"] == 1
        assert result["status"] == "completed"

        client.post("/api/demo/reset")


def test_auto_run_requires_enabled_preference():
    with TestClient(app) as client:
        client.post("/api/demo/reset")
        client.put("/api/preferences", json={"auto_submit_enabled": False})
        response = client.post("/api/applications/auto-run", json={"max_items": 1})
        assert response.status_code == 400


def test_discard_hides_restore_shows():
    with TestClient(app) as client:
        client.post("/api/demo/reset")
        job_id = client.get("/api/jobs").json()["items"][0]["id"]

        discarded = client.delete(f"/api/jobs/{job_id}")
        assert discarded.status_code == 200
        assert discarded.json()["job"]["status"] == "discarded"

        normal_ids = {item["id"] for item in client.get("/api/jobs").json()["items"]}
        assert job_id not in normal_ids
        discarded_ids = {item["id"] for item in client.get("/api/jobs?discarded=true").json()["items"]}
        assert job_id in discarded_ids

        restored = client.post(f"/api/jobs/{job_id}/restore")
        assert restored.status_code == 200
        assert restored.json()["job"]["status"] == "new"

        normal_ids = {item["id"] for item in client.get("/api/jobs").json()["items"]}
        assert job_id in normal_ids
        client.post("/api/demo/reset")


def test_freshness_filter_and_unknown_dates():
    with TestClient(app) as client:
        client.post("/api/demo/reset")
        job_ids = [item["id"] for item in client.get("/api/jobs").json()["items"][:2]]
        old_id, unknown_id = job_ids
        with SessionLocal() as db:
            old_job = db.get(Job, old_id)
            unknown_job = db.get(Job, unknown_id)
            assert old_job is not None and unknown_job is not None
            old_job.posted_at = datetime.now(UTC) - timedelta(days=10)
            unknown_job.posted_at = None
            db.commit()

        filtered_payload = client.get("/api/jobs?max_age_days=2").json()
        filtered_ids = {item["id"] for item in filtered_payload["items"]}
        assert old_id not in filtered_ids
        assert unknown_id not in filtered_ids
        assert filtered_payload["unknown_date_count"] >= 1

        with_unknown_payload = client.get("/api/jobs?max_age_days=2&include_unknown_date=true").json()
        with_unknown = {item["id"] for item in with_unknown_payload["items"]}
        assert unknown_id in with_unknown
        assert old_id not in with_unknown
        assert with_unknown_payload["unknown_date_count"] == 0

        all_ids = {item["id"] for item in client.get("/api/jobs").json()["items"]}
        assert old_id in all_ids
        assert unknown_id in all_ids
        client.post("/api/demo/reset")


def test_original_cv_endpoint_returns_uploaded_file():
    with TestClient(app) as client:
        client.post("/api/demo/reset")
        document = Document()
        document.add_paragraph("Ana Torres")
        document.add_paragraph("Backend Developer")
        document.add_paragraph("ana@example.com")
        buffer = BytesIO()
        document.save(buffer)
        content = buffer.getvalue()

        uploaded = client.post(
            "/api/cv/upload",
            files={
                "file": (
                    "ana.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert uploaded.status_code == 200
        payload = uploaded.json()
        assert payload["profile"]["has_cv"] is True
        suggestions = payload["search_suggestions"]
        assert "Backend Developer" in suggestions["roles"]
        assert "Desarrollador Backend" in suggestions["roles"]
        prefs = client.get("/api/preferences").json()["preferences"]
        assert "Backend Developer" in prefs["roles"]

        original = client.get("/api/cv/original")
        assert original.status_code == 200
        assert original.content == content

        client.post("/api/demo/reset")


def test_manual_import_bypasses_freshness_filter():
    with TestClient(app) as client:
        client.post("/api/demo/reset")
        job_id = client.get("/api/jobs").json()["items"][0]["id"]
        with SessionLocal() as db:
            job = db.get(Job, job_id)
            assert job is not None
            job.posted_at = datetime.now(UTC) - timedelta(days=30)
            job.imported_manually = True
            db.commit()

        filtered = client.get("/api/jobs?max_age_days=2").json()["items"]
        assert any(item["id"] == job_id and item["imported_manually"] is True for item in filtered)
        client.post("/api/demo/reset")


def test_reimport_preserves_discarded_status():
    from app.services.search import normalize_job_data, upsert_job

    with TestClient(app) as client:
        client.post("/api/demo/reset")
        job_id = client.get("/api/jobs").json()["items"][0]["id"]
        client.delete(f"/api/jobs/{job_id}")
        with SessionLocal() as db:
            job = db.get(Job, job_id)
            assert job is not None and job.status == "discarded"
            data = normalize_job_data(job.raw_json or {})
            _, created = upsert_job(db, data)
            assert created is False
            db.refresh(job)
            assert job.status == "discarded"
        client.post("/api/demo/reset")


def test_jobs_api_can_hide_or_show_stretch_matches():
    with TestClient(app) as client:
        client.post("/api/demo/reset")
        with SessionLocal() as db:
            job = Job(
                source="test",
                external_id="stretch-list-test",
                title="Mid Backend Engineer",
                company="Test Co",
                location_text="Bogotá, Colombia",
                city="Bogotá",
                country="Colombia",
                remote_type="onsite",
                remote_scope="national",
                seniority="mid",
                years_experience=2,
                posted_at=datetime.now(UTC),
                hash="stretch-list-test-hash",
                status="new",
            )
            db.add(job)
            db.flush()
            db.add(
                Match(
                    job=job,
                    score=60,
                    eligible=True,
                    subscores={},
                    reasons=["Oportunidad potencial"],
                    gaps=[],
                    red_flags=[],
                    recommendation="Oportunidad potencial",
                    model="heuristic",
                    fit_level="stretch",
                    stretch_reasons=["Nivel mid vs junior"],
                )
            )
            db.commit()
            job_id = job.id

        hidden = client.get("/api/jobs?include_stretch=false").json()["items"]
        assert all(item["id"] != job_id for item in hidden)
        visible = client.get("/api/jobs?include_stretch=true").json()["items"]
        assert any(item["id"] == job_id and item["match"]["fit_level"] == "stretch" for item in visible)

        with SessionLocal() as db:
            job = db.get(Job, job_id)
            if job is not None:
                db.delete(job)
                db.commit()
        client.post("/api/demo/reset")
