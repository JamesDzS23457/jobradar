import logging
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from app.core.config import Settings, get_settings
from app.models import Job, Match, Profile, TailoredCV
from app.services.cv_parser import ALIASES
from app.services.llm import llm_service

logger = logging.getLogger(__name__)


def _canonical(value: str) -> str:
    return ALIASES.get(value.strip().lower(), value.strip())


def _matched_skills(profile: Profile, job: Job) -> list[str]:
    profile_skills = {_canonical(skill).lower() for skill in (profile.skills or [])}
    matched = [skill for skill in (job.skills_required or []) if _canonical(skill).lower() in profile_skills]
    matched += [skill for skill in (job.skills_nice or []) if _canonical(skill).lower() in profile_skills]
    return list(dict.fromkeys(matched))


def _sort_bullets(bullets: list[str], keywords: set[str]) -> list[str]:
    def score(bullet: str) -> int:
        normalized = bullet.lower()
        return sum(1 for keyword in keywords if keyword.lower() in normalized)

    return sorted(bullets, key=score, reverse=True)


def fallback_tailored(profile: Profile, job: Job, match: Match | None) -> dict[str, Any]:
    matched = _matched_skills(profile, job)
    keywords = {skill.lower() for skill in [*(job.skills_required or []), *(job.skills_nice or [])]}
    experience: list[dict[str, Any]] = []
    for item in profile.experience or []:
        bullets = [str(bullet) for bullet in item.get("bullets", []) if str(bullet).strip()]
        experience.append(
            {
                "company": item.get("company", ""),
                "role": item.get("role", ""),
                "dates": item.get("dates", ""),
                "bullets": _sort_bullets(bullets, keywords),
                "source_bullets": list(range(len(bullets))),
            }
        )

    strengths = list(match.reasons) if match else []
    gaps = list(match.gaps) if match else []
    priority = list(dict.fromkeys([*matched, *(profile.skills or [])]))[:14]
    summary = profile.summary.strip()
    if job.title and job.company:
        summary = f"{summary} Me interesa especialmente la posición de {job.title} en {job.company}."

    cover_letter = (
        f"Estimado equipo de {job.company}:\n\n"
        f"Me interesa la posición de {job.title}. Soy {profile.headline or 'desarrolladora backend'} "
        f"con {profile.years_experience} años de experiencia construyendo productos con "
        f"{', '.join(priority[:5])}. En mi experiencia reciente en "
        f"{profile.experience[0].get('company', 'mi empresa actual') if profile.experience else 'mi equipo'}, "
        "trabajé en APIs, bases de datos y despliegues en la nube, colaborando con producto y frontend.\n\n"
        f"Creo que puedo aportar en {job.company} por mi experiencia con {', '.join(matched[:4]) or 'tecnologías relevantes'}. "
        "Quedo atenta a una conversación para profundizar en cómo puedo contribuir al equipo.\n\n"
        f"Saludos,\n{profile.name}"
    )

    warnings = []
    if not llm_service.configured:
        warnings.append(
            "Generado con heurísticas locales. Revisa el resultado y personaliza el resumen antes de enviarlo."
        )
    if gaps:
        warnings.append("Hay brechas frente a la vacante; no inventes experiencia para cubrirlas.")

    return {
        "headline": f"{job.title} · {', '.join(priority[:3])}" if priority else job.title,
        "summary": summary,
        "skills_priority": priority,
        "experience": experience,
        "education": profile.education or [],
        "keywords_added": matched,
        "cover_letter": cover_letter,
        "strengths": strengths,
        "gaps": gaps,
        "warnings": warnings,
    }


def _merge_llm_content(content: dict[str, Any], fallback: dict[str, Any]) -> dict[str, Any]:
    merged = {**fallback, **{key: value for key, value in content.items() if value not in (None, "", [], {})}}
    if not merged.get("experience"):
        merged["experience"] = fallback["experience"]
    if not merged.get("summary"):
        merged["summary"] = fallback["summary"]
    if not merged.get("cover_letter"):
        merged["cover_letter"] = fallback["cover_letter"]
    merged.setdefault("warnings", [])
    return merged


def generate_tailored(db, job: Job, profile: Profile, match: Match | None, settings: Settings | None = None) -> TailoredCV:
    settings = settings or get_settings()
    fallback = fallback_tailored(profile, job, match)
    content = fallback

    if llm_service.configured:
        profile_payload = profile.as_dict()
        job_payload = {
            "title": job.title,
            "company": job.company,
            "description": job.description,
            "skills_required": job.skills_required,
            "skills_nice": job.skills_nice,
            "seniority": job.seniority,
            "languages": job.languages,
            "location": job.location_text,
        }
        llm_content = llm_service.tailor_cv(profile_payload, job_payload)
        if llm_content:
            content = _merge_llm_content(llm_content, fallback)
            content["model"] = settings.llm_model_quality
        else:
            content["warnings"] = [*content.get("warnings", []), "El LLM no respondió; se usó el generador local."]
    else:
        content["model"] = "heuristic"

    tailored = TailoredCV(job_id=job.id, content=content)
    db.add(tailored)
    db.commit()
    db.refresh(tailored)

    try:
        tailored.docx_path = str(export_docx(tailored, profile, settings))
        tailored.pdf_path = str(export_pdf(tailored, profile, settings))
        db.commit()
        db.refresh(tailored)
    except Exception:  # pragma: no cover - file system dependent
        logger.exception("Could not export tailored CV files")

    return tailored


def _contact_line(profile: Profile) -> str:
    parts = [profile.email, profile.phone, profile.location]
    return " · ".join(part for part in parts if part)


def export_docx(tailored: TailoredCV, profile: Profile, settings: Settings | None = None) -> Path:
    settings = settings or get_settings()
    content = tailored.content or {}
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    name = document.add_paragraph()
    name.paragraph_format.space_after = Pt(2)
    run = name.add_run(profile.name or "Tu nombre")
    run.bold = True
    run.font.size = Pt(22)

    contact = document.add_paragraph(_contact_line(profile))
    contact.paragraph_format.space_after = Pt(8)

    headline = document.add_paragraph(content.get("headline", ""))
    headline.paragraph_format.space_after = Pt(10)
    for run in headline.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = None

    def heading(text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = None

    if content.get("summary"):
        heading("Perfil")
        document.add_paragraph(content["summary"])

    if content.get("skills_priority"):
        heading("Habilidades")
        document.add_paragraph(" · ".join(content["skills_priority"]))

    if content.get("experience"):
        heading("Experiencia")
        for item in content["experience"]:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(f"{item.get('role', '')} — {item.get('company', '')}")
            run.bold = True
            if item.get("dates"):
                dates = paragraph.add_run(f"  |  {item['dates']}")
                dates.italic = True
            for bullet in item.get("bullets", []):
                document.add_paragraph(str(bullet), style="List Bullet")

    if content.get("education"):
        heading("Educación")
        for item in content["education"]:
            document.add_paragraph(f"{item.get('degree', '')} — {item.get('institution', '')} ({item.get('dates', '')})")

    path = settings.export_dir / f"cv_{tailored.id}_{_slug(profile.name) or 'perfil'}.docx"
    document.save(path)
    return path


def export_pdf(tailored: TailoredCV, profile: Profile, settings: Settings | None = None) -> Path:
    settings = settings or get_settings()
    content = tailored.content or {}
    path = settings.export_dir / f"cv_{tailored.id}_{_slug(profile.name) or 'perfil'}.pdf"

    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "Name",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=23,
        alignment=TA_LEFT,
        spaceAfter=2,
    )
    contact_style = ParagraphStyle(
        "Contact", parent=styles["Normal"], fontSize=9, leading=12, textColor="#475569", spaceAfter=8
    )
    headline_style = ParagraphStyle(
        "Headline", parent=styles["Normal"], fontSize=10.5, leading=13, textColor="#0f766e", spaceAfter=8
    )
    heading_style = ParagraphStyle(
        "Heading", parent=styles["Heading2"], fontSize=10, leading=12, spaceBefore=8, spaceAfter=4, textColor="#0f172a"
    )
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=9.5, leading=13, spaceAfter=2)
    bullet_style = ParagraphStyle("Bullet", parent=body_style, leftIndent=10, bulletIndent=2)

    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
        title=f"CV {profile.name}",
        author=profile.name,
    )
    story: list[Any] = [
        Paragraph(escape(profile.name or "Tu nombre"), name_style),
        Paragraph(escape(_contact_line(profile)), contact_style),
        Paragraph(escape(str(content.get("headline", ""))), headline_style),
    ]

    def add_heading(text: str) -> None:
        story.append(Spacer(1, 4))
        story.append(Paragraph(escape(text.upper()), heading_style))

    if content.get("summary"):
        add_heading("Perfil")
        story.append(Paragraph(escape(str(content["summary"])).replace("\n", "<br/>"), body_style))

    if content.get("skills_priority"):
        add_heading("Habilidades")
        story.append(Paragraph(escape(" · ".join(content["skills_priority"])), body_style))

    if content.get("experience"):
        add_heading("Experiencia")
        for item in content["experience"]:
            story.append(
                Paragraph(
                    f"<b>{escape(str(item.get('role', '')))}</b> — {escape(str(item.get('company', '')))}"
                    f" <font color='#64748b'>{escape(str(item.get('dates', '')))}</font>",
                    body_style,
                )
            )
            bullets = [str(bullet) for bullet in item.get("bullets", []) if str(bullet).strip()]
            if bullets:
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(escape(bullet), bullet_style), leftIndent=12) for bullet in bullets],
                        bulletType="bullet",
                        start="•",
                        leftIndent=12,
                    )
                )

    if content.get("education"):
        add_heading("Educación")
        for item in content["education"]:
            story.append(
                Paragraph(
                    f"{escape(str(item.get('degree', '')))} — {escape(str(item.get('institution', '')))} "
                    f"({escape(str(item.get('dates', '')))})",
                    body_style,
                )
            )

    doc.build(story)
    return path


def _slug(value: str) -> str:
    return "".join(char for char in (value or "").lower().replace(" ", "_") if char.isalnum() or char in "_-")[:40]


def tailored_to_dict(tailored: TailoredCV) -> dict[str, Any]:
    return {
        "id": tailored.id,
        "job_id": tailored.job_id,
        "content": tailored.content,
        "approved": tailored.approved,
        "created_at": tailored.created_at,
        "docx_url": f"/api/tailored/{tailored.id}/download?format=docx",
        "pdf_url": f"/api/tailored/{tailored.id}/download?format=pdf",
    }
